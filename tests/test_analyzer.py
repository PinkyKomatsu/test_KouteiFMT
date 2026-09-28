"""フォーマット解析のテスト（受け入れテスト 2）。"""
from __future__ import annotations

import openpyxl
import pytest
from openpyxl.styles import Border, Side

from app import analyzer
from app import xlsx_writer as xw

BOX = Border(*(Side(style="thin"),) * 4)


def by_name(fmt, name):
    return next(f for f in fmt["fields"] if f["name"] == name)


def test_fields_detected(fmt):
    expect = {
        "報告日": ("date", "right", ["B3"]),
        "工程": ("phase", "right", ["D3"]),
        "進捗": ("status", "right", ["F3"]),
        "報告者": ("reporter", "right", ["B4"]),
        "件名": ("subject", "right", ["B5"]),
        "報告内容": ("body", "down", ["A8"]),
        "課題": ("issue", "right", ["B16"]),
        "今後の予定": ("plan", "right", ["B19"]),
        "No": ("list", "table", [f"A{r}" for r in range(24, 29)]),
        "アクション": ("list", "table", [f"B{r}" for r in range(24, 29)]),
        "担当": ("list", "table", [f"D{r}" for r in range(24, 29)]),
        "期限": ("list", "table", [f"E{r}" for r in range(24, 29)]),
        "状態": ("list", "table", [f"F{r}" for r in range(24, 29)]),
        "Summary (EN)": ("body", "right", ["B30"]),
        "上長コメント": ("skip", "right", ["B34"]),
        "備考": ("body", "right", ["B38"]),
    }
    got = {f["name"]: (f["kind"], f["direction"], f["cells"]) for f in fmt["fields"]}
    assert got == expect
    assert len({f["table_id"] for f in fmt["fields"] if f["direction"] == "table"}) == 1


def test_supervisor_comment_is_skip_and_disabled(fmt):
    f = by_name(fmt, "上長コメント")
    assert f["kind"] == "skip" and f["enabled"] is False


def test_summary_en_is_english_and_paired(fmt):
    en = by_name(fmt, "Summary (EN)")
    main = by_name(fmt, "報告内容")
    assert en["lang"] == "en"
    assert en["partner_of"] == main["id"] and main["pair"] == en["id"]
    assert all(f["lang"] == "ja" for f in fmt["fields"] if f["name"] != "Summary (EN)")


def test_title_row_below_is_not_a_field(fmt):
    assert not any(f["label_cell"] == "A1" for f in fmt["fields"])
    assert not any(xw.split_ref(r)[1] == 2 for f in fmt["fields"] for r in f["cells"])


def test_ambiguous_placement_needs_review(fmt):
    note = by_name(fmt, "備考")
    assert note["review"] is True and note["confidence"] < analyzer.REVIEW_THRESHOLD
    assert {c["direction"] for c in note["candidates"]} == {"right", "down"}
    assert any("右と下" in r for r in note["reasons"])
    # はっきりした欄は要確認にならない
    for name in ("報告日", "件名", "課題", "今後の予定", "アクション", "Summary (EN)"):
        assert not by_name(fmt, name)["review"], name
    assert [f["name"] for f in analyzer.review_fields(fmt)] == ["備考"]


@pytest.fixture()
def tricky_book(tmp_path):
    """日英の対・両言語ラベル・記入先のないラベルを含む小さなフォーマット。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["C1"] = "件名", "Subject"
    for ref in ("B1", "D1"):
        ws[ref].border = BOX
    ws["A3"] = "工程 / Phase"
    ws["B3"].border = BOX
    ws["A5"] = "報告日"          # 右にも下にも枠がない → 記入先未確定
    ws["A7"] = "コメント"
    ws.merge_cells("B7:D7")      # 結合だけで枠のない欄 → 要確認
    path = tmp_path / "tricky.xlsx"
    wb.save(path)
    return path


def test_bilingual_pairs_and_unassigned(tricky_book):
    fmt = analyzer.analyze(tricky_book)
    subject, subject_en = by_name(fmt, "件名"), by_name(fmt, "Subject")
    assert subject_en["lang"] == "en" and subject_en["kind"] == "subject"
    assert subject["pair"] == subject_en["id"] and subject_en["partner_of"] == subject["id"]
    assert by_name(fmt, "工程 / Phase")["lang"] == "both"

    date_field = by_name(fmt, "報告日")
    assert date_field["cells"] == [] and date_field["review"] and date_field["direction"] == "none"

    comment = by_name(fmt, "コメント")
    assert comment["cells"] == ["B7"] and comment["review"]
    assert any("枠のない欄" in r for r in comment["reasons"])

    # ユーザーがセルを指定すると要確認が外れる
    analyzer.assign_cells(tricky_book, date_field, ["B5"])
    assert date_field["cells"] == ["B5"] and not date_field["review"] and date_field["confidence"] == 1.0


def test_field_from_dragged_range(template, fmt):
    f = analyzer.field_from_cells(template, fmt["sheets"][0], ["B24", "C24", "B25", "B26"], fmt["fields"])
    assert f["cells"] == ["B24", "B25", "B26"], "結合セルはアンカーにそろえる"
    assert f["name"].startswith("アクション") and f["kind"] == "list"


def test_sheet_layout_for_preview(template):
    lay = analyzer.sheet_layout(template)
    cells = {c["ref"]: c for c in lay["cells"]}
    assert cells["A8"]["r2"] == 15 and cells["A8"]["c2"] == 6, "結合範囲"
    assert cells["A3"]["bold"] and cells["A3"]["fill"] == "#DDE5F0" and cells["A3"]["label"]
    assert cells["B3"]["border"] and not cells["A2"]["border"]
    assert lay["col_px"][0] == 110 and lay["row_px"][0] == 43


def test_analyze_rejects_xls(samples):
    xls = next(samples["reports_dir"].rglob("*.xls"))
    with pytest.raises(xw.XlsNotSupported):
        analyzer.analyze(xls)
