"""フォーマット解析：Excel を読み、ラベルと記入欄を推定してフォーマット定義を作る。

- openpyxl は読み取りだけに使う（保存はしない）。結合セルは左上のアンカーセルに正規化する。
- 各欄に確信度（0〜1）を付け、迷うもの（確信度 0.6 未満・候補が 2 つ・枠のない欄・記入先なし）は「要確認」にする。
- 欄ごとの記入言語（日本語／英語／日英併記）を判定し、日本語欄と英語欄の対を検出する。
"""
from __future__ import annotations

import io
import re
import warnings
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl

from . import xlsx_writer as xw

KIND_LABELS = {
    "date": "日付", "subject": "件名", "phase": "工程", "status": "進捗",
    "reporter": "報告者", "dept": "部署", "body": "本文", "issue": "課題",
    "plan": "今後の予定", "request": "依頼事項", "list": "表", "skip": "記入しない",
}
DIRECTION_LABELS = {"right": "右", "down": "下", "inline": "同じセル", "table": "表", "manual": "手動", "none": "未確定"}
LANG_LABELS = {"ja": "日本語", "en": "英語", "both": "日英併記"}
TEXT_KINDS = ("body", "issue", "plan", "request")   # メモを振り分ける本文系の欄

# 上から順に判定する（英語ラベルは大文字小文字を区別しない）
KIND_RULES = [
    ("skip", ("承認", "検印", "上長", "宛先", "approv")),
    ("date", ("日付", "報告日", "作成日", "date")),
    ("subject", ("件名", "テーマ", "表題", "subject", "title")),
    ("phase", ("工程", "フェーズ", "phase")),
    ("status", ("進捗", "ステータス", "status")),
    ("reporter", ("報告者", "担当", "氏名", "reporter")),
    ("dept", ("部署", "所属", "department", "dept")),
    ("issue", ("課題", "懸念", "リスク", "issue", "risk")),
    ("plan", ("今後", "予定", "next", "plan")),
    ("request", ("依頼", "相談", "判断", "request")),
    ("body", ("本文",)),
]

PLACEHOLDER_RE = re.compile(
    r"^\s*(|〇+|○+|◯+|[xXｘＸ×]{2,}|[＿_]+|[-－]+|（ここに記入）|\(ここに記入\)|ここに記入|記入欄|未記入)\s*$"
)
INLINE_RE = re.compile(
    r"^(?P<label>.{1,20}?)\s*(?P<colon>[：:])\s*(?:〇+|○+|◯+|[xXｘＸ×]{2,}|[＿_]+|（ここに記入）|\(ここに記入\))\s*$"
)
_NAME_STRIP_RE = re.compile(r"^[\s【\[［■●◆・]+|[\s】\]］：:]+$")
_JA_RE = re.compile(r"[぀-ヿ㐀-鿿ｦ-ﾟ]")
_EN_RE = re.compile(r"[A-Za-z]{2,}")
_EN_MARK_RE = re.compile(r"[（(]\s*(EN|ENG|English|英語|英)\s*[)）]|英語|English", re.I)

REVIEW_THRESHOLD = 0.6
MAX_ROWS = 400
MAX_COLS = 60


def classify(name: str) -> str:
    low = (name or "").lower()
    for kind, words in KIND_RULES:
        if any(w in low for w in words):
            return kind
    return "body"


def label_lang(text: str) -> str:
    """ラベルの言語：'ja' / 'en' / 'both'（「件名 / Subject」のように両方ある）。"""
    t = text or ""
    ja, en = bool(_JA_RE.search(t)), bool(_EN_RE.search(t))
    if _EN_MARK_RE.search(t) and not re.search(r"[/／]", t):
        return "en"
    if ja and en:
        return "both"
    if en:
        return "en"
    return "ja"


def clean_name(text: str) -> str:
    s = _NAME_STRIP_RE.sub("", (text or "").strip())
    return s or (text or "").strip()


def is_placeholder(value) -> bool:
    if value is None:
        return True
    return isinstance(value, str) and bool(PLACEHOLDER_RE.match(value))


# ---------------------------------------------------------------------------
# 読み込み（図形・画像は解析に不要なので外してから openpyxl に渡す）
# ---------------------------------------------------------------------------

def _without_drawings(path: Path) -> io.BytesIO:
    """シートの rels から drawing の参照を外したコピーをメモリ上に作る。
    openpyxl は画像を読むのに Pillow を必要とするため（原本ファイル自体は変更しない）。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if re.match(r"^xl/worksheets/_rels/[^/]+\.rels$", info.filename):
                root = xw.parse_xml(data)
                for r in list(root):
                    if (r.get("Type") or "").endswith("/drawing"):
                        root.remove(r)
                data = xw.serialize_xml(root)
            zout.writestr(info.filename, data)
    buf.seek(0)
    return buf


def load_book(path, data_only: bool = True):
    path = Path(path)
    xw.check_supported(path)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(_without_drawings(path), data_only=data_only)


# ---------------------------------------------------------------------------
# グリッド（結合セルをアンカーに正規化）
# ---------------------------------------------------------------------------

@dataclass
class Anchor:
    r1: int
    c1: int
    r2: int
    c2: int
    value: object
    border: bool
    merged: bool

    @property
    def ref(self) -> str:
        return xw.make_ref(self.c1, self.r1)

    @property
    def text(self) -> str:
        return self.value.strip() if isinstance(self.value, str) else ""

    @property
    def empty(self) -> bool:
        return is_placeholder(self.value)

    @property
    def fillable(self) -> bool:
        """記入欄になり得る空欄（罫線も結合もない空白は余白とみなす）。"""
        return self.empty and (self.border or self.merged)


def _has_border(cell) -> bool:
    try:
        b = cell.border
    except AttributeError:
        return False
    if b is None:
        return False
    return any(getattr(getattr(b, side, None), "style", None) for side in ("left", "right", "top", "bottom"))


def col_px(ws, c: int) -> int:
    dim = ws.column_dimensions.get(xw.index_to_col(c))
    if dim is not None and dim.hidden:
        return 0
    width = dim.width if dim is not None and dim.width else None
    return int(round(width * 7 + 5)) if width else 64


def row_px(ws, r: int) -> int:
    dim = ws.row_dimensions.get(r)
    if dim is not None and dim.hidden:
        return 0
    height = dim.height if dim is not None and dim.height else None
    return int(round(height * 96 / 72)) if height else 20


class Grid:
    def __init__(self, ws):
        self.ws = ws
        self.owner: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        for mr in ws.merged_cells.ranges:
            rng = (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    self.owner[(r, c)] = rng
        keys = list(ws._cells.keys()) + list(self.owner.keys())
        self.max_row = min(max((k[0] for k in keys), default=1), MAX_ROWS)
        self.max_col = min(max((k[1] for k in keys), default=1), MAX_COLS)
        self._cache: dict[tuple[int, int], Anchor] = {}
        self.col_px = [0] + [col_px(ws, c) for c in range(1, self.max_col + 2)]
        self.row_px = [0] + [row_px(ws, r) for r in range(1, self.max_row + 2)]

    def anchor_at(self, r: int, c: int) -> Anchor:
        rng = self.owner.get((r, c))
        r1, c1, r2, c2 = rng if rng else (r, c, r, c)
        a = self._cache.get((r1, c1))
        if a is None:
            cell = self.ws._cells.get((r1, c1))
            value = cell.value if cell is not None else None
            border = any(_has_border(cl) for rr in range(r1, r2 + 1) for cc in range(c1, c2 + 1)
                         if (cl := self.ws._cells.get((rr, cc))) is not None)
            a = Anchor(r1, c1, r2, c2, value, border, rng is not None)
            self._cache[(r1, c1)] = a
        return a

    def anchor_ref(self, ref: str) -> Anchor:
        c, r = xw.split_ref(ref)
        return self.anchor_at(r, c)

    def anchors(self):
        for r in range(1, self.max_row + 1):
            for c in range(1, self.max_col + 1):
                rng = self.owner.get((r, c))
                if rng and (rng[0], rng[1]) != (r, c):
                    continue
                yield self.anchor_at(r, c)

    def width(self, a: Anchor) -> int:
        return sum(self.col_px[c] for c in range(a.c1, min(a.c2, self.max_col + 1) + 1))

    def height(self, a: Anchor) -> int:
        return sum(self.row_px[r] for r in range(a.r1, min(a.r2, self.max_row + 1) + 1))

    def area(self, a: Anchor) -> int:
        return self.width(a) * self.height(a)

    def sheet_width(self) -> int:
        return sum(self.col_px[1:self.max_col + 1])


def is_label(a: Anchor) -> bool:
    if not isinstance(a.value, str):
        return False
    s = a.value.strip()
    return 0 < len(s) <= 30 and not s.endswith("。") and not is_placeholder(s)


# ---------------------------------------------------------------------------
# 記入欄の推定
# ---------------------------------------------------------------------------

def _field(fid, name, label: Anchor | None, sheet, cells, kind, direction, grid: Grid | None,
           prefix="", table_id=None, confidence=0.9, reasons=None, candidates=None, lang=None):
    first = grid.anchor_ref(cells[0]) if (grid and cells) else None
    reasons = list(reasons or [])
    if confidence < REVIEW_THRESHOLD and not reasons:
        reasons.append(f"確信度が低い（{confidence:.2f}）")
    return {
        "id": fid,
        "name": name,
        "label_text": label.text if label else "",
        "label_cell": label.ref if label else "",
        "sheet": sheet,
        "cells": list(cells),
        "kind": kind,
        "direction": direction,
        "prefix": prefix,
        "height_px": grid.height(first) if first else 20,
        "width_px": grid.width(first) if first else 64,
        "table_id": table_id,
        "enabled": kind != "skip",
        "lang": lang or (label_lang(label.text) if label else "ja"),
        "pair": None,          # この欄と対になる英語欄の id（この欄が日本語側のとき）
        "partner_of": None,    # 対になる日本語欄の id（この欄が英語側のとき）
        "confidence": round(confidence, 2),
        "review": bool(reasons),
        "reasons": reasons,
        "candidates": candidates or [],
    }


def _detect_sheet(ws, grid: Grid, next_id, table_counter) -> list[dict]:
    fields: list[dict] = []
    used: set[tuple[int, int]] = set()
    labels = [a for a in grid.anchors() if is_label(a)]

    def key(a: Anchor):
        return (a.r1, a.c1)

    # 1. 表：同じ行に隣接するラベルが 3 つ以上、その下が枠付きの空欄
    by_row: dict[int, list[Anchor]] = {}
    for a in labels:
        by_row.setdefault(a.r1, []).append(a)
    for _, row_labels in sorted(by_row.items()):
        row_labels.sort(key=lambda a: a.c1)
        runs, run = [], [row_labels[0]]
        for a in row_labels[1:]:
            if a.c1 == run[-1].c2 + 1 and a.r2 == run[-1].r2:
                run.append(a)
            else:
                runs.append(run)
                run = [a]
        runs.append(run)
        for run in runs:
            if len(run) < 3:
                continue
            columns = []
            for h in run:
                cells, rr = [], h.r2 + 1
                while rr <= grid.max_row and len(cells) < 200:
                    b = grid.anchor_at(rr, h.c1)
                    if b.c1 != h.c1 or b.r1 != rr or not b.empty or not b.border:
                        break
                    cells.append(b)
                    rr = b.r2 + 1
                columns.append(cells)
            n = min(len(c) for c in columns)
            if n == 0:
                continue
            table_counter[0] += 1
            tid = f"t{table_counter[0]}"
            langs = [label_lang(h.text) for h in run]
            table_lang = "en" if langs.count("en") > len(langs) / 2 else "ja"
            confidence = 0.9 if all(len(c) == n for c in columns) else 0.7
            for h, cells in zip(run, columns):
                fields.append(_field(next_id(), clean_name(h.text), h, ws.title, [b.ref for b in cells[:n]],
                                     "list", "table", grid, table_id=tid, confidence=confidence, lang=table_lang))
                used.add(key(h))
                used.update(key(b) for b in cells[:n])

    # 2. 同じセル内に記入：「件名：〇〇」
    for a in labels:
        if key(a) in used:
            continue
        m = INLINE_RE.match(a.text)
        if m:
            prefix = a.text[: m.end("colon")]
            name = clean_name(m.group("label"))
            fields.append(_field(next_id(), name, a, ws.title, [a.ref], classify(name), "inline", grid,
                                 prefix=prefix, confidence=0.9, lang=label_lang(name)))
            used.add(key(a))

    # 3. 右か下の空欄：まず全ラベルの「右」を確定させ、その後で「下」を割り当てる
    pending = [a for a in labels if key(a) not in used]
    sheet_w = grid.sheet_width() or 1

    def right_of(a: Anchor):
        if a.c2 + 1 > grid.max_col:
            return None
        b = grid.anchor_at(a.r1, a.c2 + 1)
        if b.c1 == a.c2 + 1 and b.r1 == a.r1 and b.fillable and key(b) not in used:
            return b
        return None

    def down_of(a: Anchor):
        if a.r2 + 1 > grid.max_row:
            return None
        b = grid.anchor_at(a.r2 + 1, a.c1)
        if b.c1 == a.c1 and b.r1 == a.r2 + 1 and b.empty and b.border and key(b) not in used:
            return b
        return None

    chosen: dict[tuple[int, int], dict] = {}
    for a in pending:                                   # パス 1：右
        right, down = right_of(a), down_of(a)
        if right is None:
            continue
        wide = grid.width(a) >= sheet_w * 0.6
        big_down = bool(down and not right.merged and down.merged and grid.area(down) >= grid.area(right) * 2)
        if wide or big_down:
            continue                                    # 下を採る（パス 2）
        reasons, conf = [], 0.9 if right.border else 0.5
        if not right.border:
            reasons.append("枠のない欄を記入先にしました")
        cands = [{"direction": "right", "cells": [right.ref]}]
        if down:
            conf = min(conf, 0.55)
            reasons.append("右と下の両方に記入欄の候補があります")
            cands.append({"direction": "down", "cells": [down.ref]})
        chosen[key(a)] = {"label": a, "direction": "right", "target": right, "conf": conf,
                          "reasons": reasons, "cands": cands}
        used.add(key(right))
    for a in pending:                                   # パス 2：下
        if key(a) in chosen:
            continue
        right, down = right_of(a), down_of(a)
        if down is None and right is None:
            continue
        if down is None:                                # 幅の広い見出しで下がない → 右
            chosen[key(a)] = {"label": a, "direction": "right", "target": right, "conf": 0.6,
                              "reasons": [], "cands": [{"direction": "right", "cells": [right.ref]}]}
            used.add(key(right))
            continue
        cands = [{"direction": "down", "cells": [down.ref]}]
        reasons, conf = [], 0.85
        if grid.width(a) >= sheet_w * 0.6:
            conf = 0.8
        if right:
            conf = 0.55
            reasons.append("右と下の両方に記入欄の候補があります")
            cands.append({"direction": "right", "cells": [right.ref]})
        chosen[key(a)] = {"label": a, "direction": "down", "target": down, "conf": conf,
                          "reasons": reasons, "cands": cands}
        used.add(key(down))

    for c in chosen.values():
        a = c["label"]
        name = clean_name(a.text)
        fields.append(_field(next_id(), name, a, ws.title, [c["target"].ref], classify(name), c["direction"],
                             grid, confidence=c["conf"], reasons=c["reasons"],
                             candidates=c["cands"] if len(c["cands"]) > 1 else []))
        used.add(key(a))

    # 4. 記入欄らしいラベルなのに記入先が見つからない → 記入先未確定の欄として残す（ユーザーが指定する）
    for a in pending:
        if key(a) in used or key(a) in chosen:
            continue
        name = clean_name(a.text)
        kind = classify(name)
        if kind == "body" or len(name) > 12 or grid.width(a) >= sheet_w * 0.6:
            continue
        fields.append(_field(next_id(), name, a, ws.title, [], kind, "none", grid, confidence=0.0,
                             reasons=["記入先のセルが見つかりません。プレビューでセルを指定してください"]))
    return fields


def _sort_key(sheet_order, field):
    ref = field["cells"][0] if field["cells"] else (field.get("label_cell") or "A1")
    c, r = xw.split_ref(ref)
    return (sheet_order.get(field["sheet"], 0), r, c)


def _dedupe_names(fields: list[dict]) -> None:
    seen: dict[str, int] = {}
    for f in fields:
        base = f["name"]
        if base in seen:
            seen[base] += 1
            f["name"] = f"{base}({seen[base]})"
        else:
            seen[base] = 1


def main_body(fields: list[dict]) -> dict | None:
    """メイン本文欄：ラベルに「内容／概要／報告」を含む本文欄、なければ高さが最大の本文欄（日本語側）。"""
    bodies = [f for f in fields if f.get("enabled") and f["kind"] == "body" and not f.get("partner_of")
              and f.get("lang") != "en" and f.get("cells")]
    for f in bodies:
        if any(k in f["name"] for k in ("内容", "概要", "報告")):
            return f
    return max(bodies, key=lambda f: f.get("height_px", 0), default=None)


def detect_pairs(fields: list[dict]) -> None:
    """英語欄を、対になる日本語欄に結びつける（pair / partner_of）。"""
    for f in fields:
        f["pair"] = None
        f["partner_of"] = None
    singles = [f for f in fields if f["direction"] != "table" and f.get("cells")]
    ja_fields = [f for f in singles if f.get("lang") in ("ja", "both")]
    for en in [f for f in singles if f.get("lang") == "en"]:
        ec, er = xw.split_ref(en.get("label_cell") or en["cells"][0])

        def dist(j):
            jc, jr = xw.split_ref(j.get("label_cell") or j["cells"][0])
            return abs(jc - ec) + abs(jr - er)

        same_kind = [j for j in ja_fields if j["kind"] == en["kind"] and not j["pair"] and j["sheet"] == en["sheet"]]
        target = None
        if en["kind"] != "body" and same_kind:
            target = min(same_kind, key=dist)
        elif en["kind"] == "body":
            m = main_body(fields)
            if m and not m["pair"]:
                target = m
        if target is not None:
            target["pair"] = en["id"]
            en["partner_of"] = target["id"]


def analyze(path, name: str | None = None) -> dict:
    """Excel を解析してフォーマット定義（dict）を返す。"""
    path = Path(path)
    wb = load_book(path)
    counter = [0]
    table_counter = [0]

    def next_id():
        counter[0] += 1
        return f"f{counter[0]}"

    fields: list[dict] = []
    sheets = []
    for ws in wb.worksheets:
        sheets.append(ws.title)
        fields.extend(_detect_sheet(ws, Grid(ws), next_id, table_counter))
    order = {s: i for i, s in enumerate(sheets)}
    fields.sort(key=lambda f: _sort_key(order, f))
    _dedupe_names(fields)
    detect_pairs(fields)
    return {
        "version": 2,
        "name": name or path.stem,
        "template": "template" + path.suffix.lower(),
        "source": str(path),
        "created": datetime.now().isoformat(timespec="seconds"),
        "sheets": sheets,
        "fields": fields,
    }


def review_fields(fmt: dict) -> list[dict]:
    return [f for f in fmt["fields"] if f.get("enabled") and (f.get("review") or not f.get("cells"))]


# ---------------------------------------------------------------------------
# ユーザーによる指定（R7）
# ---------------------------------------------------------------------------

def new_field_id(fields: list[dict]) -> str:
    nums = [int(f["id"][1:]) for f in fields if re.match(r"^f\d+$", f.get("id", ""))]
    return f"f{max(nums, default=0) + 1}"


def _label_near(grid: Grid, target: Anchor):
    """ラベルを左 → 上の順に探す。"""
    for c in range(target.c1 - 1, max(0, target.c1 - 4), -1):
        a = grid.anchor_at(target.r1, c)
        if is_label(a):
            return a
        if not a.empty:
            break
    for r in range(target.r1 - 1, max(0, target.r1 - 4), -1):
        a = grid.anchor_at(r, target.c1)
        if is_label(a):
            return a
        if not a.empty:
            break
    return None


def anchor_refs(template, sheet: str, refs: list[str]) -> list[str]:
    """選択されたセルをアンカーに正規化し、重複を除いて行・列の順に並べる。"""
    grid = Grid(load_book(template)[sheet])
    out = []
    for r in refs:
        a = grid.anchor_ref(r)
        if a.ref not in out:
            out.append(a.ref)
    return sorted(out, key=lambda r: (xw.split_ref(r)[1], xw.split_ref(r)[0]))


def field_from_cells(template, sheet: str, refs: list[str], fields: list[dict]) -> dict:
    """選択セル（ドラッグした範囲なら表の列）を記入欄にする。"""
    grid = Grid(load_book(template)[sheet])
    cells = anchor_refs(template, sheet, refs)
    target = grid.anchor_ref(cells[0])
    label = _label_near(grid, target)
    name = clean_name(label.text) if label else target.ref
    direction = "manual"
    field = _field(new_field_id(fields), name, label, sheet, cells,
                   "list" if len(cells) > 1 else classify(name), direction, grid, confidence=1.0)
    names = {f["name"] for f in fields}
    base, n = field["name"], 1
    while field["name"] in names:
        n += 1
        field["name"] = f"{base}({n})"
    return field


def assign_cells(template, field: dict, refs: list[str]) -> None:
    """既存の欄の記入先を、ユーザーが指定したセルに変える（要確認を解除する）。"""
    field["cells"] = anchor_refs(template, field["sheet"], refs)
    field["direction"] = "manual"
    field["confidence"] = 1.0
    field["review"] = False
    field["reasons"] = []
    field["candidates"] = []


# ---------------------------------------------------------------------------
# プレビュー用
# ---------------------------------------------------------------------------

def display_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return f"{value.year}/{value.month}/{value.day}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _fill_hex(cell) -> str | None:
    try:
        fill = cell.fill
        if fill is None or fill.fill_type != "solid":
            return None
        rgb = fill.fgColor.rgb if fill.fgColor is not None else None
        if isinstance(rgb, str) and len(rgb) in (6, 8) and rgb not in ("00000000",):
            return "#" + rgb[-6:]
    except (AttributeError, TypeError):
        pass
    return None


def sheet_layout(template, sheet: str | None = None, max_rows: int = 120, max_cols: int = 40) -> dict:
    """プレビューでシートを再現するための情報（列幅・行高・結合・背景色・太字・罫線・非表示）。"""
    wb = load_book(template)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    grid = Grid(ws)
    n_rows = min(grid.max_row, max_rows)
    n_cols = min(grid.max_col, max_cols)
    cells = []
    for a in grid.anchors():
        if a.r1 > n_rows or a.c1 > n_cols:
            continue
        cell = ws._cells.get((a.r1, a.c1))
        font = getattr(cell, "font", None)
        cells.append({
            "ref": a.ref, "r1": a.r1, "c1": a.c1,
            "r2": min(a.r2, n_rows), "c2": min(a.c2, n_cols),
            "text": display_text(a.value), "border": a.border, "label": is_label(a),
            "fill": _fill_hex(cell) if cell is not None else None,
            "bold": bool(font and font.b),
            "size": float(font.sz) if font is not None and font.sz else 11.0,
        })
    return {
        "sheet": ws.title,
        "sheets": wb.sheetnames,
        "n_rows": n_rows,
        "n_cols": n_cols,
        "col_px": grid.col_px[1:n_cols + 1],
        "row_px": grid.row_px[1:n_rows + 1],
        "cells": cells,
    }
