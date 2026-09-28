"""テスト用：架空の工程報告フォーマットと過去報告書 5 件を作る。

    python tests/make_samples.py [出力フォルダ]     （既定: tests/samples）

フォーマット（シート「工程報告書」）:
  タイトル / 報告日・工程・進捗 / 報告者 / 件名 / 【報告内容】の大枠 / 課題 / 今後の予定 /
  No・アクション・担当・期限・状態の表 / Summary (EN) / 上長コメント / 備考（右と下の両方に枠＝曖昧な配置）/
  ロゴ画像・図形・数式（calcChain 付き）・印刷設定
過去報告書のうち 1 件（移行リハーサル第 1 回）は、先頭に 1 行多い旧版（行が 1 行ずれている）。
1 件（総合テスト）には、個人情報検出の確認用にダミーの患者 ID を含める。
"""
from __future__ import annotations

import struct
import sys
import zipfile
import zlib
from datetime import datetime
from pathlib import Path

import openpyxl
from lxml import etree
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

SHEET = "工程報告書"
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GRAY = PatternFill("solid", fgColor="DDE5F0")
TOP_LEFT = Alignment(horizontal="left", vertical="top")  # 折り返しなし（書き込み時にスタイルが複製されるか確認する）

OPENING = "お疲れ様です。"
CLOSING = "以上、ご報告いたします。"
OPENING_EN = "Please find my report below."
CLOSING_EN = "That concludes my report."
DUMMY_PATIENT_ID = "1234567"

CASES = [
    {
        "file": "2026-06-05_要件定義.xlsx",
        "date": datetime(2026, 6, 5), "phase": "要件定義", "status": "完了", "reporter": "山田 太郎",
        "subject": "医事会計システム更新 要件定義の完了報告",
        "body": [OPENING,
                 "【背景】", "・現行の医事会計システムが保守期限を迎えるため、更新プロジェクトを進めています。",
                 "【状況】", "・各部門へのヒアリングを実施し、要件定義書を作成しました。", "・窓口会計と未収金管理の要件を確定しました。",
                 "【対応】", "・基本設計に着手します。",
                 CLOSING],
        "summary": [OPENING_EN,
                    "[Background]", "・The current medical accounting system is reaching the end of its maintenance period, so we are proceeding with the replacement project.",
                    "[Status]", "・We held hearings with each department and prepared the requirements definition document.", "・We finalized the requirements for front-desk billing and outstanding patient payments.",
                    "[Actions]", "・We will start the basic design.",
                    CLOSING_EN],
        "issue": ["・電子カルテ連携の仕様確認に時間がかかっています。"],
        "plan": ["・7/1から基本設計を開始する予定です。"],
        "table": [("1", "基本設計の体制確定", "山田", "6/20", "対応中")],
    },
    {
        "file": "2026-07-03_基本設計.xlsx",
        "date": datetime(2026, 7, 3), "phase": "基本設計", "status": "予定どおり", "reporter": "山田 太郎",
        "subject": "医事会計システム更新 基本設計の進捗報告",
        "body": [OPENING,
                 "【背景】", "・6月に確定した要件定義にもとづき、基本設計を進めています。",
                 "【状況】", "・画面設計と帳票設計のレビューが完了しました。", "・点数マスタの移行方針を決定しました。",
                 "【対応】", "・インターフェース設計を進めます。",
                 CLOSING],
        "summary": [OPENING_EN,
                    "[Background]", "・We are proceeding with the basic design based on the requirements finalized in June.",
                    "[Status]", "・The review of the screen and form designs has been completed.", "・We have decided on the migration policy for the fee schedule master.",
                    "[Actions]", "・We will proceed with the interface design.",
                    CLOSING_EN],
        "issue": ["・DPC関連の帳票要件が一部未確定です。"],
        "plan": ["・7/31までに基本設計書を承認いただく予定です。"],
        "table": [("1", "インターフェース設計", "佐藤", "7/20", "対応中"), ("2", "帳票要件の確定", "山田", "7/15", "未着手")],
    },
    {
        "file": "2026-08-07_開発.xlsx",
        "date": datetime(2026, 8, 7), "phase": "開発", "status": "遅延", "reporter": "山田 太郎",
        "subject": "医事会計システム更新 開発工程の進捗報告",
        "body": [OPENING,
                 "【背景】", "・基本設計の承認を受け、ベンダーによる開発を進めています。",
                 "【状況】", "・算定ロジックの開発が計画より1週間遅れています。", "・その他の機能は予定どおり進んでいます。",
                 "【対応】", "・ベンダーに要員の追加を依頼しました。",
                 CLOSING],
        "summary": [OPENING_EN,
                    "[Background]", "・Following the approval of the basic design, the vendor is proceeding with development.",
                    "[Status]", "・The development of the fee calculation logic is one week behind schedule.", "・Other functions are progressing as planned.",
                    "[Actions]", "・We have asked the vendor to add staff.",
                    CLOSING_EN],
        "issue": ["・算定ロジックの遅れが総合テストに影響する懸念があります。"],
        "plan": ["・8/20までに遅れを解消する予定です。"],
        "table": [("1", "要員追加の調整", "山田", "8/12", "対応中")],
    },
    {
        "file": "2026-09-04_総合テスト.xlsx",
        "date": datetime(2026, 9, 4), "phase": "総合テスト", "status": "予定どおり", "reporter": "佐藤 花子",
        "subject": "医事会計システム更新 総合テストの結果報告",
        "body": [OPENING,
                 "【背景】", "・開発が完了し、総合テストを実施しました。",
                 "【状況】", "・テストケースの98%が合格しました。", "・レセプト出力で2件の不具合が見つかりました。",
                 "【対応】", "・不具合の修正版を9/10に適用します。",
                 CLOSING],
        "summary": [OPENING_EN,
                    "[Background]", "・Development has been completed, and we conducted the system integration test.",
                    "[Status]", "・98% of the test cases passed.", "・Two defects were found in the medical claim (receipt) output.",
                    "[Actions]", "・We will apply the fix on 9/10.",
                    CLOSING_EN],
        "issue": [f"・検証データに患者ID {DUMMY_PATIENT_ID} が含まれていたため、匿名化データに差し替えました。"],
        "plan": ["・9/15から移行リハーサルを開始する予定です。"],
        "table": [("1", "不具合の修正", "佐藤", "9/10", "対応中"), ("2", "移行リハーサルの準備", "山田", "9/12", "未着手")],
    },
    {
        "file": "2026-09-18_移行リハーサル_旧版.xlsx",
        "shift": 1,
        "date": datetime(2026, 9, 18), "phase": "移行リハーサル", "status": "予定どおり", "reporter": "山田 太郎",
        "subject": "医事会計システム更新 移行リハーサル（第1回）の実施報告",
        "body": [OPENING,
                 "【背景】", "・本番移行に向けて、移行リハーサル（第1回）を実施しました。",
                 "【状況】", "・データ移行は予定の時間内に完了しました。", "・移行後のレセプト点検で差異は見つかりませんでした。",
                 "【対応】", "・第2回の移行リハーサルで切り戻し手順を確認します。",
                 CLOSING],
        "summary": [OPENING_EN,
                    "[Background]", "・We conducted the first migration rehearsal in preparation for the production cutover.",
                    "[Status]", "・The data migration was completed within the planned time.", "・No discrepancies were found in the medical claim check after the migration.",
                    "[Actions]", "・We will confirm the rollback procedure in the second migration rehearsal.",
                    CLOSING_EN],
        "issue": ["・切り戻し手順が未確認です。"],
        "plan": ["・9/26に移行リハーサル（第2回）を実施する予定です。"],
        "table": [("1", "切り戻し手順の作成", "佐藤", "9/24", "対応中"), ("2", "稼働判定基準の作成", "山田", "9/30", "未着手")],
    },
]


# ---------------------------------------------------------------------------
# シートの組み立て
# ---------------------------------------------------------------------------

def _box(ws, rng: str):
    if ":" in rng and rng not in [str(r) for r in ws.merged_cells.ranges]:
        ws.merge_cells(rng)
    cells = ws[rng]
    if not isinstance(cells, tuple):
        cells = ((cells,),)
    for row in cells:
        for c in (row if isinstance(row, tuple) else (row,)):
            c.border = BOX


def _label(ws, rng: str, text: str):
    _box(ws, rng)
    c = ws[rng.split(":")[0]]
    c.value = text
    c.fill = GRAY
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


# 各欄の行（shift で 1 行ずらす）
ROWS = {"title": 1, "row3": 3, "row4": 4, "subject": 5, "body_head": 7, "body": (8, 15), "issue": (16, 18),
        "plan": (19, 21), "table_head": 23, "table": (24, 28), "summary": (30, 32), "comment": (34, 36),
        "note": 38, "note_below": (39, 40), "formula": 42}


def build_sheet(ws, values: dict | None = None, off: int = 0):
    v = values or {}

    def R(r: int) -> int:
        return r + off

    for col in "ABCDEF":
        ws.column_dimensions[col].width = 15
    if off:
        ws["A1"] = "社外秘"
    ws.merge_cells(f"A{R(1)}:F{R(1)}")
    ws[f"A{R(1)}"] = "医事会計システム更新 工程報告書"
    ws[f"A{R(1)}"].font = Font(bold=True, size=16)
    ws[f"A{R(1)}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[R(1)].height = 32

    _label(ws, f"A{R(3)}", "報告日")
    _box(ws, f"B{R(3)}")
    _label(ws, f"C{R(3)}", "工程")
    _box(ws, f"D{R(3)}")
    _label(ws, f"E{R(3)}", "進捗")
    _box(ws, f"F{R(3)}")
    _label(ws, f"A{R(4)}", "報告者")
    _box(ws, f"B{R(4)}")
    _label(ws, f"A{R(5)}", "件名")
    _box(ws, f"B{R(5)}:F{R(5)}")

    _label(ws, f"A{R(7)}:F{R(7)}", "【報告内容】")
    _box(ws, f"A{R(8)}:F{R(15)}")
    ws[f"A{R(8)}"].alignment = TOP_LEFT
    _label(ws, f"A{R(16)}:A{R(18)}", "課題")
    _box(ws, f"B{R(16)}:F{R(18)}")
    ws[f"B{R(16)}"].alignment = TOP_LEFT
    _label(ws, f"A{R(19)}:A{R(21)}", "今後の予定")
    _box(ws, f"B{R(19)}:F{R(21)}")
    ws[f"B{R(19)}"].alignment = TOP_LEFT

    _label(ws, f"A{R(23)}", "No")
    _label(ws, f"B{R(23)}:C{R(23)}", "アクション")
    _label(ws, f"D{R(23)}", "担当")
    _label(ws, f"E{R(23)}", "期限")
    _label(ws, f"F{R(23)}", "状態")
    for r in range(R(24), R(29)):
        for rng in (f"A{r}", f"B{r}:C{r}", f"D{r}", f"E{r}", f"F{r}"):
            _box(ws, rng)

    _label(ws, f"A{R(30)}:A{R(32)}", "Summary (EN)")
    _box(ws, f"B{R(30)}:F{R(32)}")
    ws[f"B{R(30)}"].alignment = TOP_LEFT
    _label(ws, f"A{R(34)}:A{R(36)}", "上長コメント")
    _box(ws, f"B{R(34)}:F{R(36)}")
    # 曖昧な配置：右にも下にも枠付きの空欄がある
    _label(ws, f"A{R(38)}", "備考")
    _box(ws, f"B{R(38)}:F{R(38)}")
    _box(ws, f"A{R(39)}:F{R(40)}")

    ws[f"A{R(42)}"] = f"=COUNTA(B{R(24)}:B{R(28)})"

    ws.print_area = f"A1:F{R(40)}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    if v:
        ws[f"B{R(3)}"] = v["date"]
        ws[f"B{R(3)}"].number_format = "yyyy/mm/dd"
        ws[f"D{R(3)}"] = v["phase"]
        ws[f"F{R(3)}"] = v["status"]
        ws[f"B{R(4)}"] = v["reporter"]
        ws[f"B{R(5)}"] = v["subject"]
        ws[f"A{R(8)}"] = "\n".join(v["body"])
        ws[f"B{R(16)}"] = "\n".join(v["issue"])
        ws[f"B{R(19)}"] = "\n".join(v["plan"])
        ws[f"B{R(30)}"] = "\n".join(v["summary"])
        for i, row in enumerate(v["table"]):
            r = R(24) + i
            for col, val in zip(("A", "B", "D", "E", "F"), row):
                ws[f"{col}{r}"] = val


# ---------------------------------------------------------------------------
# ロゴ画像・図形・calcChain を ZIP に直接追加する（openpyxl の画像追加は Pillow が必要なため）
# ---------------------------------------------------------------------------

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"

DRAWING_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<xdr:oneCellAnchor>
<xdr:from><xdr:col>5</xdr:col><xdr:colOff>76200</xdr:colOff><xdr:row>0</xdr:row><xdr:rowOff>38100</xdr:rowOff></xdr:from>
<xdr:ext cx="838200" cy="304800"/>
<xdr:pic>
<xdr:nvPicPr><xdr:cNvPr id="2" name="Logo"/><xdr:cNvPicPr><a:picLocks noChangeAspect="1"/></xdr:cNvPicPr></xdr:nvPicPr>
<xdr:blipFill><a:blip r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>
<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="838200" cy="304800"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>
</xdr:pic>
<xdr:clientData/>
</xdr:oneCellAnchor>
<xdr:twoCellAnchor>
<xdr:from><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>40</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
<xdr:to><xdr:col>6</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>42</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="3" name="Stamp"/><xdr:cNvSpPr/></xdr:nvSpPr>
<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln w="9525"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill></a:ln></xdr:spPr>
</xdr:sp>
<xdr:clientData/>
</xdr:twoCellAnchor>
</xdr:wsDr>
"""

DRAWING_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/></Relationships>
"""

CALC_CHAIN = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="A42" i="1"/></calcChain>
"""

_BEFORE_DRAWING = ("printOptions", "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks",
                   "customProperties", "cellWatches", "ignoredErrors", "smartTags")


def png_bytes(w: int, h: int, rgb=(30, 70, 140)) -> bytes:
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _dump(root) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def add_drawing_and_calcchain(path: Path) -> None:
    with zipfile.ZipFile(path) as z:
        order = [i.filename for i in z.infolist()]
        parts = {n: z.read(n) for n in order}

    sheet = "xl/worksheets/sheet1.xml"
    parts["xl/media/image1.png"] = png_bytes(110, 40)
    parts["xl/drawings/drawing1.xml"] = DRAWING_XML.encode("utf-8")
    parts["xl/drawings/_rels/drawing1.xml.rels"] = DRAWING_RELS.encode("utf-8")

    rels_name = "xl/worksheets/_rels/sheet1.xml.rels"
    rels = (etree.fromstring(parts[rels_name]) if rels_name in parts
            else etree.Element("{%s}Relationships" % NS_PKG, nsmap={None: NS_PKG}))
    etree.SubElement(rels, "{%s}Relationship" % NS_PKG, Id="rIdLogo1",
                     Type=NS_R + "/drawing", Target="../drawings/drawing1.xml")
    parts[rels_name] = _dump(rels)

    root = etree.fromstring(parts[sheet])
    drawing = etree.Element("{%s}drawing" % NS_MAIN, nsmap={"r": NS_R})
    drawing.set("{%s}id" % NS_R, "rIdLogo1")
    last = None
    for child in root:
        if etree.QName(child).localname in _BEFORE_DRAWING:
            last = child
    if last is not None:
        last.addnext(drawing)
    else:
        root.append(drawing)
    parts[sheet] = _dump(root)

    parts["xl/calcChain.xml"] = CALC_CHAIN.encode("utf-8")
    wb_rels = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
    etree.SubElement(wb_rels, "{%s}Relationship" % NS_PKG, Id="rIdCalc1", Type=NS_R + "/calcChain", Target="calcChain.xml")
    parts["xl/_rels/workbook.xml.rels"] = _dump(wb_rels)

    ct = etree.fromstring(parts["[Content_Types].xml"])
    if not any(d.get("Extension") == "png" for d in ct.iter("{%s}Default" % NS_CT)):
        first_override = ct.find("{%s}Override" % NS_CT)
        png = etree.Element("{%s}Default" % NS_CT, Extension="png", ContentType="image/png")
        if first_override is not None:
            first_override.addprevious(png)
        else:
            ct.append(png)
    etree.SubElement(ct, "{%s}Override" % NS_CT, PartName="/xl/drawings/drawing1.xml",
                     ContentType="application/vnd.openxmlformats-officedocument.drawing+xml")
    etree.SubElement(ct, "{%s}Override" % NS_CT, PartName="/xl/calcChain.xml",
                     ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml")
    parts["[Content_Types].xml"] = _dump(ct)

    for name in parts:
        if name not in order:
            order.append(name)
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name in order:
            z.writestr(name, parts[name])
    tmp.replace(path)


def _new_book(values=None, off=0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    build_sheet(ws, values, off)
    return wb


def make_all(out_dir) -> dict:
    out = Path(out_dir)
    reports = out / "reports"
    reports.mkdir(parents=True, exist_ok=True)

    fmt_path = out / "format.xlsx"
    _new_book().save(fmt_path)
    add_drawing_and_calcchain(fmt_path)

    paths = []
    for case in CASES:
        p = reports / case["file"]
        _new_book(case, case.get("shift", 0)).save(p)
        paths.append(p)

    sub = reports / "old"
    sub.mkdir(exist_ok=True)
    (reports / f"~${CASES[0]['file']}").write_bytes(b"lock file")
    (sub / "旧形式の報告書.xls").write_bytes(b"\xd0\xcf\x11\xe0dummy")
    return {"format": fmt_path, "reports_dir": reports, "reports": paths, "cases": CASES}


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "samples"
    info = make_all(target)
    print(f"フォーマット: {info['format']}")
    for p in info["reports"]:
        print(f"過去報告書: {p}")
